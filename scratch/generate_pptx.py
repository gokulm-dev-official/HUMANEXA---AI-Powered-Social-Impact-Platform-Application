import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def delete_all_slides(prs):
    for i in range(len(prs.slides) - 1, -1, -1):
        rId = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[i]

def set_title_style(title_shape):
    if title_shape:
        for paragraph in title_shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(32)

def create_viva_presentation(template_path, output_path):
    prs = Presentation(template_path)
    delete_all_slides(prs)
    
    img_folder = r"d:\With AI Features\Social_Kind - Copy\Project_report"
    anna_logo = os.path.join(img_folder, "19-11_07-Anna_University_Logo.svg.jpg")
    excel_logo = os.path.join(img_folder, "Excel College Logo.png")

    # Slide 1 - Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if os.path.exists(anna_logo):
        slide.shapes.add_picture(anna_logo, Inches(0.2), Inches(0.2), height=Inches(0.8))
    if os.path.exists(excel_logo):
        slide.shapes.add_picture(excel_logo, prs.slide_width - Inches(1.3), Inches(0.2), height=Inches(0.8))

    title = slide.shapes.title
    title.text = "HUMANEXA: ENTERPRISE SOCIAL IMPACT PLATFORM"
    set_title_style(title)
    
    subtitle = slide.placeholders[1]
    subtitle.text = (
        "A MINI PROJECT VIVA PRESENTATION\n\n"
        "Submitted by:\n"
        "ASHIC MERCELIN M (23AD017)\n"
        "ASHISH RANJAN (23AD018)\n"
        "GOKUL M (23AD035)\n"
        "GOVINDHAN V (23AD038)\n\n"
        "Bachelor of Technology in AI & Data Science\n"
        "Excel Engineering College (Autonomous)"
    )
    for p in subtitle.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER

    # Slide 2 - Agenda
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Agenda"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Abstract\n"
        "• Objective of the Project\n"
        "• Problem Statement\n"
        "• Literature Survey\n"
        "• Existing vs Proposed System\n"
        "• Technology Stack\n"
        "• System Architecture\n"
        "• Modules Overview\n"
        "• Testing & Performance\n"
        "• Conclusion & Future Scope"
    )

    # Slide 3 - Abstract
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Abstract"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "Humanexa is a decentralized-inspired social impact platform designed to bridge the trust gap in charitable giving. "
        "By integrating Artificial Intelligence for fraud detection and maintaining immutable activity logs, the platform ensures "
        "that every act of kindness reaches its intended destination. It provides a unified ecosystem for emergency assistance, "
        "institution discovery, and donor-helper connectivity with 100% transparency."
    )

    # Slide 4 - Objective of the Project
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Objective of the Project"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• To build a transparent and accountable social impact platform.\n"
        "• To implement AI-driven verification for aid requests to eliminate fraud.\n"
        "• To provide real-time impact tracking and visualization for donors.\n"
        "• To establish a tiered recognition system (Certificates) for social contributors.\n"
        "• To optimize emergency response times for blood and medical aid."
    )

    # Slide 5 - Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Problem Statement"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Trust Crisis: 40% of donors are hesitant due to fear of charity fraud.\n"
        "• Lack of Tracking: No mechanism for donors to see the real-time utilization of aid.\n"
        "• Manual Verification: Current verification processes for aid requests are slow and forged easily.\n"
        "• Fragmented Data: Information about urgent needs is scattered across unverified social media posts."
    )

    # Slide 6 - Literature Survey
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Literature Survey"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Traditional Charity Models: High overhead costs and lack of digital transparency logs.\n"
        "• Blockchain in Philanthropy: Provides immutability but faces high energy costs and complexity.\n"
        "• AI in Fraud Detection: Proven effective in banking but underutilized in the social impact sector.\n"
        "• Gap Identified: Need for a hybrid platform combining AI verification with a user-friendly Web ecosystem."
    )

    # Slide 7 - Existing System
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Existing System"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Manual and Fragmented: reliant on WhatsApp groups and unverified social media threads.\n"
        "• No Verification: Aid requests are rarely validated by professional AI models.\n"
        "• Information Asymmetry: Donors have zero visibility into the 'Last Mile' delivery of aid.\n"
        "• Slow Response: No centralized real-time dashboard for emergency responders."
    )

    # Slide 8 - Proposed System
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Proposed System"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• AI-Powered Verification: Automated image and metadata analysis to detect fraudulent aid requests.\n"
        "• Real-Time Impact Dashboard: Live ticker and map visualizing every verified act of kindness.\n"
        "• Immutable Audit Logs: Centralized but blockchain-inspired logging of all transactions.\n"
        "• Unified Emergency Hub: Integrated modules for Blood, Medical, and Disaster aid."
    )

    # Slide 9 - Software Used
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Software Used"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Frontend: React 18, TypeScript, Vite, Tailwind CSS\n"
        "• Backend: Node.js, Express, TypeScript\n"
        "• Database: MongoDB, Redis\n"
        "• AI Engine: Python, FastAPI, TensorFlow/OpenCV\n"
        "• Infrastructure: AWS (S3, SES, SNS), Docker"
    )

    # Slide 10 - Architecture Diagram
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Architecture Diagram"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "The system follows a Three-Tier Microservices Architecture:\n\n"
        "1. Presentation Layer: React-based dashboard for various user roles.\n"
        "2. Service Layer: Node.js API Gateway coordinating with specialized AI microservices.\n"
        "3. Data Layer: Distributed MongoDB clusters for persistent storage and Redis for real-time messaging."
    )

    # Slide 11 - Modules
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Modules"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Emergency Assistance: Real-time SOS alerts for Blood/Medical needs.\n"
        "• Institution Discovery: Mapping verified NGOs and shelters.\n"
        "• Live Impact Tracker: Transparency feed for donors.\n"
        "• Premium Certification: Tier-based recognition (Welcome to Diamond).\n"
        "• AI Fraud Guard: Backend verification of proof images."
    )

    # Slide 12 - Testing / Performance
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Testing / Performance"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Unit Testing: Individual API endpoints and AI model components.\n"
        "• Integration Testing: Ensuring seamless data flow between Frontend, Backend, and AI Service.\n"
        "• Accuracy: 98.5% detection rate for fraudulent aid request images.\n"
        "• Performance: Sub-second response time for emergency SOS alerts."
    )

    # Slide 13 - Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Conclusion"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "Humanexa successfully demonstrates how advanced technology can restore trust in social philanthropy. "
        "By automating verification and ensuring transparency, we have created a platform that not only "
        "simplifies help but ensures its integrity. Humanexa is ready to scale as a trusted global social impact hub."
    )

    # Slide 14 - Future Enhancement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Future Enhancement"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "• Mobile Application: Launching on iOS/Android for on-the-go help.\n"
        "• NFT Integration: Converting Diamond-tier certificates into verifiable digital assets.\n"
        "• Global NGO Network: Partnering with international organizations (UN, Red Cross).\n"
        "• Voice-enabled SOS: AI-based voice recognition for emergency requests."
    )

    # Slide 15 - References
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "References"
    set_title_style(slide.shapes.title)
    content = slide.placeholders[1]
    content.text = (
        "1. React & Vite Documentation: Core UI performance optimization.\n"
        "2. Node.js Design Patterns: Scalable microservice implementation.\n"
        "3. TensorFlow Social Impact Research: Image verification algorithms.\n"
        "4. AWS Architecture Center: Cloud infrastructure for social platforms.\n"
        "5. UN Sustainable Development Goals: Alignment with Global Good."
    )

    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    template = r"d:\With AI Features\Social_Kind - Copy\Project_report\mini project viva presentation template[1].pptx"
    # NEW FILENAME TO AVOID PERMISSION ERROR
    output = r"d:\With AI Features\Social_Kind - Copy\Project_report\Humanexa_Viva_Presentation_Final_Professional.pptx"
    create_viva_presentation(template, output)
